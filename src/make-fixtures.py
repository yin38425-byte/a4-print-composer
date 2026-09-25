from pathlib import Path
import sys
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root = Path(__file__).parent
qa = Path(sys.argv[1]) if len(sys.argv)>1 else root / 'qa'
qa.mkdir(exist_ok=True)
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
d.styles['Normal'].font.name = 'SimSun'
d.styles['Normal'].font.size = Pt(12)
d.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
d.add_heading('多格式混合排版测试', 0)
d.add_paragraph('这是软件测试样张，不含真实票据或个人资料。')
d.add_paragraph('第一页：中文、英文 Mixed formats、数字 1234.56。')
t=d.add_table(rows=1, cols=2)
t.style='Table Grid'
t.rows[0].cells[0].text='检查项目'
t.rows[0].cells[1].text='预期内容'
for left,right in [('中文表格','内容和边框保留'),('分页','显式分页后进入第二页')]:
    cells=t.add_row().cells; cells[0].text=left; cells[1].text=right
d.add_picture(str(root.parent/'qa'/'chinese-image.png'),width=Mm(130))
d.add_page_break()
d.add_heading('第二页原始内容',1)
d.add_paragraph('尾页核查标记 END-OF-DOCX-2026。此行不得丢失。')
d.save(qa/'mixed-chinese.docx')
(qa/'simple.rtf').write_text(r'{\rtf1\ansi\deff0 {\fonttbl{\f0 Arial;}}\f0\fs28 RTF mixed format sample\par FINAL-RTF-MARKER}',encoding='ascii')
(qa/'broken.docx').write_bytes(b'not a word document')
print(qa)

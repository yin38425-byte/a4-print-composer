"""Synthetic presentation fixtures; no real user documents or external media."""
import io
import struct
import sys
import zipfile
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)

def chunk(tag, data):
    return tag + struct.pack('<I', len(data)) + data + (b'\0' if len(data) % 2 else b'')

# Two seconds of actual MJPEG video. Blue first frame differs from the red poster.
frames = []
for color in ['#1565c0', '#1565c0', '#43a047', '#43a047']:
    buf = io.BytesIO()
    Image.new('RGB', (320, 180), color).save(buf, 'JPEG')
    frames.append(buf.getvalue())
avih = struct.pack('<14I', 500000, 0, 0, 16, 4, 0, 1, max(map(len, frames)), 320, 180, 0, 0, 0, 0)
strh = struct.pack('<4s4sIHHIIIIIIIIhhhh', b'vids', b'MJPG', 0, 0, 0, 0, 1, 2, 0, 4, max(map(len, frames)), 0xffffffff, 0, 0, 0, 320, 180)
strf = struct.pack('<IiiHH4sIiiII', 40, 320, 180, 1, 24, b'MJPG', 320*180*3, 0, 0, 0, 0)
hdrl = chunk(b'LIST', b'hdrl' + chunk(b'avih', avih) + chunk(b'LIST', b'strl' + chunk(b'strh', strh) + chunk(b'strf', strf)))
movie = b''
index = b''
for frame in frames:
    index += struct.pack('<4sIII', b'00dc', 16, len(movie)+4, len(frame))
    movie += chunk(b'00dc', frame)
(out/'synthetic-video.avi').write_bytes(chunk(b'RIFF', b'AVI ' + hdrl + chunk(b'LIST', b'movi'+movie) + chunk(b'idx1', index)))
poster = Image.new('RGB', (640, 360), '#dd3344')
draw = ImageDraw.Draw(poster)
draw.rectangle((320, 0, 639, 359), fill='#00bacc')
draw.text((35, 160), 'VIDEO POSTER - NOT FIRST FRAME', fill='white', font_size=24)
poster.save(out/'video-poster.png')

def slide(prs, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(.6), Inches(.3), Inches(11), Inches(.7))
    box.text_frame.paragraphs[0].text = title
    box.text_frame.paragraphs[0].font.size = Pt(28)
    return s

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(12), Inches(6.75)
s = slide(prs, 'PPTX-1 中文演示：多格式混合排版')
s.shapes.add_picture(str(out/'video-poster.png'), Inches(1), Inches(1.8), width=Inches(4))
s = slide(prs, 'PPTX-2 视频封面验证：红色与青色')
s.shapes.add_movie(str(out/'synthetic-video.avi'), Inches(2), Inches(1.5), Inches(8), Inches(4.5), poster_frame_image=str(out/'video-poster.png'), mime_type='video/avi')
s = slide(prs, 'PPTX-3 隐藏页也应按原顺序导入')
s._element.set('show', '0')
prs.save(out/'synthetic-video.pptx')
legacy = Presentation()
slide(legacy, 'PPT-LEGACY 中文旧格式测试')
legacy.save(out/'legacy-input.pptx')
(out/'broken.pptx').write_text('This is not a presentation', encoding='utf-8')
with zipfile.ZipFile(out/'synthetic-video.pptx') as src:
    for filename, extra in [('macro.pptx', ('ppt/vbaProject.bin', b'synthetic macro marker')), ('external.pptx', ('ppt/slides/_rels/external.xml.rels', b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/video" TargetMode="External" Target="https://example.invalid/video.avi"/></Relationships>'))]:
        with zipfile.ZipFile(out/filename, 'w', zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                dst.writestr(item, src.read(item.filename))
            dst.writestr(*extra)
print('Created synthetic PPT fixtures in', out)
